from __future__ import annotations

import csv
import json
import sys
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.worksheet import Worksheet


# ---------------------------------------------------------------------------
# STYLE CONSTANTS
# ---------------------------------------------------------------------------

DARK_BLUE = "1F4E78"
MEDIUM_BLUE = "5B9BD5"
LIGHT_BLUE = "D9EAF7"
LIGHT_GREEN = "E2F0D9"
LIGHT_ORANGE = "FCE4D6"
LIGHT_RED = "F4CCCC"
LIGHT_GRAY = "E7E6E6"
DARK_GRAY = "666666"
WHITE = "FFFFFF"
BLACK = "000000"
GREEN_TEXT = "008000"
ORANGE_TEXT = "C65911"

THIN_GRAY_BORDER = Border(
    bottom=Side(
        style="thin",
        color="D9E1F2",
    )
)


# ---------------------------------------------------------------------------
# FILE HELPERS
# ---------------------------------------------------------------------------

def load_json_file(
    file_path: Path,
    *,
    default: Any = None,
) -> Any:
    if not file_path.exists():
        return default

    with file_path.open(
        "r",
        encoding="utf-8",
    ) as file:
        return json.load(file)


def load_csv_file(
    file_path: Path,
) -> list[dict[str, str]]:
    if not file_path.exists():
        return []

    with file_path.open(
        "r",
        encoding="utf-8-sig",
        newline="",
    ) as file:
        return list(
            csv.DictReader(file)
        )


def load_text_file(
    file_path: Path,
) -> str:
    if not file_path.exists():
        return ""

    return file_path.read_text(
        encoding="utf-8",
    )


# ---------------------------------------------------------------------------
# GENERIC EXCEL HELPERS
# ---------------------------------------------------------------------------

def hide_gridlines(
    worksheet: Worksheet,
) -> None:
    worksheet.sheet_view.showGridLines = False


def style_title(
    worksheet: Worksheet,
    title: str,
    *,
    end_column: int,
) -> None:
    worksheet.merge_cells(
        start_row=1,
        start_column=1,
        end_row=1,
        end_column=end_column,
    )

    cell = worksheet.cell(
        row=1,
        column=1,
        value=title,
    )

    cell.font = Font(
        bold=True,
        color=WHITE,
        size=15,
    )

    cell.fill = PatternFill(
        fill_type="solid",
        fgColor=DARK_BLUE,
    )

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[1].height = 26


def style_table_header(
    worksheet: Worksheet,
    row_number: int,
    columns: int,
) -> None:
    for column_number in range(
        1,
        columns + 1,
    ):
        cell = worksheet.cell(
            row=row_number,
            column=column_number,
        )

        cell.font = Font(
            bold=True,
            color=WHITE,
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor=MEDIUM_BLUE,
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )


def apply_data_alignment(
    cell: Cell,
) -> None:
    cell.alignment = Alignment(
        vertical="top",
        wrap_text=True,
    )

    cell.border = THIN_GRAY_BORDER


def add_hyperlink(
    cell: Cell,
    url: str,
    *,
    label: str | None = None,
) -> None:
    clean_url = str(
        url or ""
    ).strip()

    if not clean_url:
        cell.value = ""
        return

    cell.value = (
        label
        if label is not None
        else clean_url
    )

    cell.hyperlink = clean_url

    cell.font = Font(
        color="0563C1",
        underline="single",
    )


def autosize_columns(
    worksheet: Worksheet,
    *,
    maximum_width: int = 45,
    minimum_width: int = 10,
) -> None:
    for column_cells in worksheet.columns:
        column_letter = get_column_letter(
            column_cells[0].column
        )

        maximum_length = 0

        for cell in column_cells:
            value = cell.value

            if value is None:
                continue

            text = str(value)

            longest_line = max(
                (
                    len(line)
                    for line in text.splitlines()
                ),
                default=0,
            )

            maximum_length = max(
                maximum_length,
                longest_line,
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(
                maximum_length + 2,
                minimum_width,
            ),
            maximum_width,
        )


def add_excel_table(
    worksheet: Worksheet,
    *,
    start_row: int,
    end_row: int,
    end_column: int,
    table_name: str,
) -> None:
    if end_row <= start_row:
        return

    table_reference = (
        f"A{start_row}:"
        f"{get_column_letter(end_column)}"
        f"{end_row}"
    )

    table = Table(
        displayName=table_name,
        ref=table_reference,
    )

    style = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )

    table.tableStyleInfo = style

    worksheet.add_table(
        table
    )


def write_records_sheet(
    worksheet: Worksheet,
    *,
    title: str,
    records: list[dict[str, Any]],
    columns: list[str],
    table_name: str,
    hyperlink_columns: dict[str, str] | None = None,
    status_column: str | None = None,
) -> None:
    hide_gridlines(
        worksheet
    )

    style_title(
        worksheet,
        title,
        end_column=max(
            1,
            len(columns),
        ),
    )

    header_row = 3

    for column_number, column_name in enumerate(
        columns,
        start=1,
    ):
        worksheet.cell(
            row=header_row,
            column=column_number,
            value=column_name,
        )

    style_table_header(
        worksheet,
        header_row,
        len(columns),
    )

    hyperlink_columns = (
        hyperlink_columns
        or {}
    )

    for row_number, record in enumerate(
        records,
        start=header_row + 1,
    ):
        for column_number, column_name in enumerate(
            columns,
            start=1,
        ):
            cell = worksheet.cell(
                row=row_number,
                column=column_number,
            )

            value = record.get(
                column_name,
                "",
            )

            if column_name in hyperlink_columns:
                add_hyperlink(
                    cell,
                    str(value or ""),
                    label=hyperlink_columns[
                        column_name
                    ],
                )

            else:
                cell.value = value

            apply_data_alignment(
                cell
            )

            if (
                status_column
                and column_name
                == status_column
            ):
                status = str(
                    value or ""
                ).strip().upper()

                if status == "READY":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor=LIGHT_GREEN,
                    )

                elif status == "REVIEW":
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor=LIGHT_ORANGE,
                    )

                elif status in {
                    "ERROR",
                    "EXCLUDED",
                }:
                    cell.fill = PatternFill(
                        fill_type="solid",
                        fgColor=LIGHT_RED,
                    )

    last_row = (
        header_row
        + max(
            1,
            len(records),
        )
    )

    if records:
        add_excel_table(
            worksheet,
            start_row=header_row,
            end_row=header_row + len(records),
            end_column=len(columns),
            table_name=table_name,
        )

    else:
        worksheet.cell(
            row=header_row + 1,
            column=1,
            value="No records.",
        )

    worksheet.freeze_panes = "A4"

    worksheet.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(len(columns))}"
        f"{last_row}"
    )

    autosize_columns(
        worksheet
    )


# ---------------------------------------------------------------------------
# SUMMARY SHEET
# ---------------------------------------------------------------------------

def build_summary_sheet(
    worksheet: Worksheet,
    *,
    summary: dict[str, Any],
) -> None:
    """
    Build a simple operational Home sheet.

    The sheet focuses on:
    - what is ready;
    - what requires attention;
    - what the user should do next.
    """
    hide_gridlines(
        worksheet
    )

    style_title(
        worksheet,
        "NPS Weekly Package",
        end_column=7,
    )

    weekly_run_id = summary.get(
        "weekly_run_id",
        "",
    )

    window_start = summary.get(
        "window_start",
        "",
    )

    window_end = summary.get(
        "window_end",
        "",
    )

    ready_records = summary.get(
        "ready_records",
        0,
    )

    manual_review_records = summary.get(
        "manual_review_records",
        0,
    )

    clients = summary.get(
        "clients",
        0,
    )

    duplicate_issues = summary.get(
        "duplicate_issues",
        0,
    )

    # ------------------------------------------------------------------
    # WEEK INFORMATION
    # ------------------------------------------------------------------

    worksheet["A3"] = "Weekly Run"
    worksheet["B3"] = weekly_run_id

    worksheet["A4"] = "Weekly Window"
    worksheet["B4"] = (
        f"{window_start} to {window_end}"
    )

    for cell_reference in [
        "A3",
        "A4",
    ]:
        worksheet[
            cell_reference
        ].font = Font(
            bold=True,
            color=DARK_GRAY,
        )

    # ------------------------------------------------------------------
    # OPERATIONAL METRICS
    # ------------------------------------------------------------------

    worksheet.merge_cells(
        start_row=6,
        start_column=1,
        end_row=6,
        end_column=7,
    )

    worksheet["A6"] = "OPERATIONAL STATUS"

    worksheet["A6"].font = Font(
        bold=True,
        color=WHITE,
        size=11,
    )

    worksheet["A6"].fill = PatternFill(
        fill_type="solid",
        fgColor=DARK_BLUE,
    )

    worksheet["A6"].alignment = Alignment(
        horizontal="left",
        vertical="center",
    )

    worksheet.row_dimensions[6].height = 22

    metric_cards = [
        {
            "start_column": 1,
            "end_column": 2,
            "label": "READY TO CREATE",
            "value": ready_records,
            "fill": LIGHT_GREEN,
        },
        {
            "start_column": 3,
            "end_column": 4,
            "label": "MANUAL REVIEW",
            "value": manual_review_records,
            "fill": LIGHT_ORANGE,
        },
        {
            "start_column": 5,
            "end_column": 6,
            "label": "CLIENTS",
            "value": clients,
            "fill": LIGHT_BLUE,
        },
    ]

    for card in metric_cards:
        start_column = card[
            "start_column"
        ]

        end_column = card[
            "end_column"
        ]

        worksheet.merge_cells(
            start_row=8,
            start_column=start_column,
            end_row=8,
            end_column=end_column,
        )

        worksheet.merge_cells(
            start_row=9,
            start_column=start_column,
            end_row=10,
            end_column=end_column,
        )

        label_cell = worksheet.cell(
            row=8,
            column=start_column,
            value=card["label"],
        )

        value_cell = worksheet.cell(
            row=9,
            column=start_column,
            value=card["value"],
        )

        label_cell.font = Font(
            bold=True,
            color=DARK_GRAY,
        )

        label_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        value_cell.font = Font(
            bold=True,
            size=24,
        )

        value_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        for row_number in range(
            8,
            11,
        ):
            for column_number in range(
                start_column,
                end_column + 1,
            ):
                worksheet.cell(
                    row=row_number,
                    column=column_number,
                ).fill = PatternFill(
                    fill_type="solid",
                    fgColor=card["fill"],
                )

    # ------------------------------------------------------------------
    # NEXT ACTIONS
    # ------------------------------------------------------------------

    worksheet.merge_cells(
        start_row=13,
        start_column=1,
        end_row=13,
        end_column=7,
    )

    worksheet["A13"] = "NEXT ACTIONS"

    worksheet["A13"].font = Font(
        bold=True,
        color=WHITE,
        size=11,
    )

    worksheet["A13"].fill = PatternFill(
        fill_type="solid",
        fgColor=DARK_BLUE,
    )

    actions = [
        (
            "1",
            "Review manual cases",
            (
                f"{manual_review_records} record(s) "
                "require confirmation."
                if manual_review_records
                else "No manual reviews required."
            ),
        ),
        (
            "2",
            "Create NPS in Salesforce",
            (
                f"Process the {ready_records} READY "
                "record(s) from the Worklist."
            ),
        ),
        (
            "3",
            "Update Slack Canvas",
            (
                "Open selam_slack_canvas.html, "
                "copy each client table, and paste it "
                "into the corresponding Slack Canvas."
            ),
        ),
    ]

    current_row = 15

    for number, action, detail in actions:
        number_cell = worksheet.cell(
            row=current_row,
            column=1,
            value=number,
        )

        action_cell = worksheet.cell(
            row=current_row,
            column=2,
            value=action,
        )

        detail_cell = worksheet.cell(
            row=current_row,
            column=3,
            value=detail,
        )

        worksheet.merge_cells(
            start_row=current_row,
            start_column=3,
            end_row=current_row,
            end_column=7,
        )

        number_cell.font = Font(
            bold=True,
            color=WHITE,
        )

        number_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=MEDIUM_BLUE,
        )

        number_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        action_cell.font = Font(
            bold=True,
        )

        detail_cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        worksheet.row_dimensions[
            current_row
        ].height = 34

        current_row += 2

    # ------------------------------------------------------------------
    # ALERTS
    # ------------------------------------------------------------------

    if duplicate_issues:
        worksheet.merge_cells(
            start_row=current_row + 1,
            start_column=1,
            end_row=current_row + 1,
            end_column=7,
        )

        alert_cell = worksheet.cell(
            row=current_row + 1,
            column=1,
            value=(
                "ATTENTION: "
                f"{duplicate_issues} duplicate issue(s) "
                "were detected. Review the Issues sheet."
            ),
        )

        alert_cell.fill = PatternFill(
            fill_type="solid",
            fgColor=LIGHT_RED,
        )

        alert_cell.font = Font(
            bold=True,
        )

        alert_cell.alignment = Alignment(
            wrap_text=True,
            vertical="center",
        )

    # ------------------------------------------------------------------
    # MANUAL REVIEW PREVIEW
    # ------------------------------------------------------------------

    review_items = summary.get(
        "manual_review_items",
        [],
    )

    review_start = current_row + 4

    if review_items:
        worksheet.merge_cells(
            start_row=review_start,
            start_column=1,
            end_row=review_start,
            end_column=7,
        )

        worksheet.cell(
            row=review_start,
            column=1,
            value="MANUAL REVIEW REQUIRED",
        )

        worksheet.cell(
            row=review_start,
            column=1,
        ).font = Font(
            bold=True,
            color=WHITE,
        )

        worksheet.cell(
            row=review_start,
            column=1,
        ).fill = PatternFill(
            fill_type="solid",
            fgColor=ORANGE_TEXT,
        )

        headers = [
            "Country",
            "Client",
            "Project",
            "Date",
            "Reason",
            "Suggested Action",
            "Salesforce",
        ]

        header_row = review_start + 1

        for column_number, header in enumerate(
            headers,
            start=1,
        ):
            worksheet.cell(
                row=header_row,
                column=column_number,
                value=header,
            )

        style_table_header(
            worksheet,
            header_row,
            len(headers),
        )

        for row_number, review in enumerate(
            review_items,
            start=header_row + 1,
        ):
            values = [
                review.get(
                    "country",
                    "",
                ),
                review.get(
                    "client",
                    "",
                ),
                review.get(
                    "project",
                    "",
                ),
                review.get(
                    "date",
                    "",
                ),
                review.get(
                    "reason",
                    "",
                ),
                review.get(
                    "suggested_action",
                    "",
                ),
                review.get(
                    "project_url",
                    "",
                ),
            ]

            for column_number, value in enumerate(
                values,
                start=1,
            ):
                cell = worksheet.cell(
                    row=row_number,
                    column=column_number,
                )

                if column_number == 7:
                    add_hyperlink(
                        cell,
                        str(value),
                        label="Open Salesforce",
                    )

                else:
                    cell.value = value

                apply_data_alignment(
                    cell
                )

    worksheet.column_dimensions[
        "A"
    ].width = 12

    worksheet.column_dimensions[
        "B"
    ].width = 28

    worksheet.column_dimensions[
        "C"
    ].width = 30

    worksheet.column_dimensions[
        "D"
    ].width = 18

    worksheet.column_dimensions[
        "E"
    ].width = 34

    worksheet.column_dimensions[
        "F"
    ].width = 42

    worksheet.column_dimensions[
        "G"
    ].width = 20

    worksheet.freeze_panes = "A3"


# ---------------------------------------------------------------------------
# SLACK READY SHEET
# ---------------------------------------------------------------------------

def split_slack_markdown_by_client(
    markdown: str,
) -> list[dict[str, str]]:
    blocks = []

    current_client = ""
    current_lines = []

    for line in markdown.splitlines():
        if line.startswith(
            "## "
        ):
            if (
                current_client
                and current_lines
            ):
                blocks.append(
                    {
                        "Client": current_client,
                        "Slack Copy Block": (
                            "\n".join(
                                current_lines
                            ).strip()
                        ),
                    }
                )

            current_client = (
                line[3:].strip()
            )

            current_lines = [
                line,
            ]

        elif current_client:
            current_lines.append(
                line
            )

    if (
        current_client
        and current_lines
    ):
        blocks.append(
            {
                "Client": current_client,
                "Slack Copy Block": (
                    "\n".join(
                        current_lines
                    ).strip()
                ),
            }
        )

    return blocks


def build_slack_sheet(
    worksheet: Worksheet,
    *,
    markdown: str,
) -> None:
    hide_gridlines(
        worksheet
    )

    style_title(
        worksheet,
        "Slack Ready — Copy one client block at a time",
        end_column=2,
    )

    worksheet["A2"] = (
        "Select the complete text in column B and copy it "
        "directly into Slack. The original Markdown file "
        "remains the canonical copy-ready output."
    )

    worksheet.merge_cells(
        start_row=2,
        start_column=1,
        end_row=2,
        end_column=2,
    )

    worksheet["A2"].alignment = Alignment(
        wrap_text=True,
        vertical="top",
    )

    blocks = split_slack_markdown_by_client(
        markdown
    )

    headers = [
        "Client",
        "Slack Copy Block",
    ]

    for column_number, header in enumerate(
        headers,
        start=1,
    ):
        worksheet.cell(
            row=4,
            column=column_number,
            value=header,
        )

    style_table_header(
        worksheet,
        4,
        2,
    )

    for row_number, block in enumerate(
        blocks,
        start=5,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=block["Client"],
        )

        copy_cell = worksheet.cell(
            row=row_number,
            column=2,
            value=block[
                "Slack Copy Block"
            ],
        )

        copy_cell.alignment = Alignment(
            wrap_text=True,
            vertical="top",
        )

        copy_cell.font = Font(
            name="Consolas",
            size=10,
        )

        worksheet.row_dimensions[
            row_number
        ].height = min(
            max(
                75,
                len(
                    block[
                        "Slack Copy Block"
                    ]
                )
                / 2.5,
            ),
            300,
        )

    if not blocks:
        worksheet.cell(
            row=5,
            column=1,
            value="No Slack blocks available.",
        )

    worksheet.column_dimensions[
        "A"
    ].width = 28

    worksheet.column_dimensions[
        "B"
    ].width = 115

    worksheet.freeze_panes = "A5"


# ---------------------------------------------------------------------------
# MANIFEST SHEET
# ---------------------------------------------------------------------------

def flatten_manifest(
    data: Any,
    *,
    prefix: str = "",
) -> list[tuple[str, str]]:
    rows = []

    if isinstance(
        data,
        dict,
    ):
        for key, value in data.items():
            new_prefix = (
                f"{prefix}.{key}"
                if prefix
                else str(key)
            )

            rows.extend(
                flatten_manifest(
                    value,
                    prefix=new_prefix,
                )
            )

    elif isinstance(
        data,
        list,
    ):
        for index, value in enumerate(
            data
        ):
            new_prefix = (
                f"{prefix}[{index}]"
            )

            rows.extend(
                flatten_manifest(
                    value,
                    prefix=new_prefix,
                )
            )

    else:
        rows.append(
            (
                prefix,
                "" if data is None else str(data),
            )
        )

    return rows


def build_manifest_sheet(
    worksheet: Worksheet,
    *,
    manifest: dict[str, Any],
) -> None:
    hide_gridlines(
        worksheet
    )

    style_title(
        worksheet,
        "Weekly Manifest",
        end_column=2,
    )

    worksheet["A3"] = "Field"
    worksheet["B3"] = "Value"

    style_table_header(
        worksheet,
        3,
        2,
    )

    flattened = flatten_manifest(
        manifest
    )

    for row_number, (
        field,
        value,
    ) in enumerate(
        flattened,
        start=4,
    ):
        worksheet.cell(
            row=row_number,
            column=1,
            value=field,
        )

        worksheet.cell(
            row=row_number,
            column=2,
            value=value,
        )

        apply_data_alignment(
            worksheet.cell(
                row=row_number,
                column=1,
            )
        )

        apply_data_alignment(
            worksheet.cell(
                row=row_number,
                column=2,
            )
        )

    worksheet.column_dimensions[
        "A"
    ].width = 58

    worksheet.column_dimensions[
        "B"
    ].width = 110

    worksheet.freeze_panes = "A4"


# ---------------------------------------------------------------------------
# ISSUES SHEET
# ---------------------------------------------------------------------------

def normalize_issue_records(
    issues: Any,
    *,
    issue_type: str,
) -> list[dict[str, Any]]:
    if not isinstance(
        issues,
        list,
    ):
        return []

    normalized = []

    for issue in issues:
        normalized.append(
            {
                "Issue Type": issue_type,
                "Details": json.dumps(
                    issue,
                    ensure_ascii=False,
                ),
            }
        )

    return normalized


# ---------------------------------------------------------------------------
# WORKBOOK GENERATION
# ---------------------------------------------------------------------------

def generate_weekly_excel(
    manifest_path: Path | str,
) -> Path:
    manifest_file = Path(
        manifest_path
    )

    if not manifest_file.exists():
        raise FileNotFoundError(
            "Weekly manifest does not exist: "
            f"{manifest_file}"
        )

    weekly_directory = (
        manifest_file.parent
    )

    manifest = load_json_file(
        manifest_file,
        default={},
    )

    summary = load_json_file(
        weekly_directory
        / "selam_weekly_summary.json",
        default={},
    )

    worklist = load_csv_file(
        weekly_directory
        / "selam_worklist.csv"
    )

    manual_review = load_csv_file(
        weekly_directory
        / "selam_manual_review.csv"
    )

    slack_markdown = load_text_file(
        weekly_directory
        / "selam_slack_ready.md"
    )

    consolidation_issues = load_json_file(
        weekly_directory
        / "selam_consolidation_issues.json",
        default=[],
    )

    manual_review_issues = load_json_file(
        weekly_directory
        / "selam_manual_review_issues.json",
        default=[],
    )

    weekly_run_id = str(
        manifest.get(
            "weekly_run_id",
            weekly_directory.name,
        )
    ).strip()

    output_path = (
        weekly_directory
        / f"{weekly_run_id}.xlsx"
    )

    workbook = Workbook()

    default_sheet = workbook.active
    workbook.remove(
        default_sheet
    )

    summary_sheet = workbook.create_sheet(
    "Home"
    )

    worklist_sheet = workbook.create_sheet(
        "Worklist"
    )

    manual_review_sheet = workbook.create_sheet(
        "Revision Manual"
    )

    manifest_sheet = workbook.create_sheet(
        "Manifest"
    )

    issues_sheet = workbook.create_sheet(
        "Issues"
    )

    build_summary_sheet(
        summary_sheet,
        summary=summary,
    )

    worklist_columns = [
        "Client",
        "Country",
        "Session",
        "Date",
        "Facilitator",
        "QR-Code",
        "Link NPS Survey",
        "Link Insights Facilitators",
        "Results",
        "Link for Results",
        "Project ID",
        "Project URL",
        "Output Status",
        "Output Issues",
    ]

    write_records_sheet(
        worklist_sheet,
        title="NPS Weekly Worklist",
        records=worklist,
        columns=worklist_columns,
        table_name="NPSWorklist",
        hyperlink_columns={
            "Link for Results": "Home",
            "Project URL": "Salesforce",
            "Link NPS Survey": "NPS",
            "Link Insights Facilitators": "Facilitators",
        },
        status_column="Output Status",
    )

    manual_review_columns = [
        "Country",
        "client_name",
        "project_name",
        "start_date",
        "status",
        "ignore_reporting",
        "decision",
        "decision_code",
        "decision_reason",
        "matched_field",
        "matched_pattern",
        "Suggested Action",
        "project_url",
    ]

    write_records_sheet(
        manual_review_sheet,
        title="Revision Manual Necesaria",
        records=manual_review,
        columns=manual_review_columns,
        table_name="NPSManualReview",
        hyperlink_columns={
            "project_url": "Open Salesforce",
        },
        status_column="decision",
    )

    build_manifest_sheet(
        manifest_sheet,
        manifest=manifest,
    )

    issue_records = (
        normalize_issue_records(
            consolidation_issues,
            issue_type=(
                "Worklist consolidation"
            ),
        )
        + normalize_issue_records(
            manual_review_issues,
            issue_type=(
                "Manual review consolidation"
            ),
        )
    )

    write_records_sheet(
        issues_sheet,
        title="Issues and Warnings",
        records=issue_records,
        columns=[
            "Issue Type",
            "Details",
        ],
        table_name="NPSIssues",
    )

    workbook.properties.title = (
        weekly_run_id
    )

    workbook.properties.subject = (
        "Weekly NPS operational package"
    )

    workbook.properties.creator = (
        "NPS Automation"
    )

    workbook.save(
        output_path
    )

    return output_path


def main() -> None:
    if len(
        sys.argv
    ) < 2:
        raise SystemExit(
            "Usage:\n"
            "py -m engine.weekly_excel "
            "<weekly_manifest_path>"
        )

    manifest_path = Path(
        sys.argv[1]
    )

    output_path = generate_weekly_excel(
        manifest_path
    )

    print("=" * 80)
    print("NPS WEEKLY EXCEL")
    print("=" * 80)
    print(f"Excel generated: {output_path}")
    print("=" * 80)


if __name__ == "__main__":
    main()